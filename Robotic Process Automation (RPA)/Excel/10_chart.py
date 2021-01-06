from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.line_chart import LineChart


wb = load_workbook("sample.xlsx")       # sample.xlsx file load
ws = wb.active

# Get data from B2 to C11
chart_value = Reference(ws, min_row=2, max_row=11, min_col=2, max_col=3)

# Create the Bar Chart (Bar, Line, Pie, etc.)
bar_chart = BarChart()
bar_chart.add_data(chart_value)     # Add the data into the bar chart
ws.add_chart(bar_chart, "E2")       # Add the bar chart at "E2" cell




# Get the data with title name
chart_value = Reference(ws, min_row=1, max_row=11, min_col=2, max_col=3)

# Create the Line Chart
line_chart = LineChart()
# Add the data into the line chart, and title name is on
line_chart.add_data(chart_value, titles_from_data=True)

line_chart.title = "Grade"          # Set the title name
line_chart.style = 20               # Apply the existed style
line_chart.x_axis.title = "ID"      # The name of X axis
line_chart.y_axis.title = "Grade"   # The name of Y axis

ws.add_chart(line_chart, "M2")      # Add the line chart at "M2" cell



wb.save("sample_chart.xlsx")
wb.close()                              # Close the work book